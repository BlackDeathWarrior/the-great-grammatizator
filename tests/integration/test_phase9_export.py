"""Phase 9: export. Structured JSON to real files.

Video is a PACKAGE, rendered deterministically. Never generative video.
"""

import os
import pathlib

import pytest

from app.export import dispatch, video
from app.formats import registry

pytestmark = pytest.mark.integration

CLAIMS = [{"text": "rated 9.1 out of 10", "chunk_id": "c1"}]

LINKEDIN = {
    "hook": "Check your version today.",
    "body": "A critical authentication bypass is being exploited in the wild.",
    "call_to_action": "Upgrade to 22.7R2.6.",
    "hashtags": ["#CyberSecurity", "#PatchNow"],
    "claims": CLAIMS,
}

EXEC_SUMMARY = {
    "title": "Critical authentication bypass",
    "bottom_line": "Patch immediately.",
    "key_points": ["Rated 9.1 on CVSS", "Exploited in the wild", "Fixed in 22.7R2.6"],
    "recommended_action": "Upgrade to 22.7R2.6.",
    "residual_risk": "Unpatched instances remain exposed.",
    "claims": CLAIMS,
}

PRESENTATION = {
    "title": "Critical bypass briefing",
    "slides": [
        {
            "heading": "What happened",
            "bullets": ["Authentication bypass", "Rated 9.1"],
            "speaker_notes": "Open by naming the product and the version.",
        },
        {
            "heading": "What to do",
            "bullets": ["Upgrade to 22.7R2.6"],
            "speaker_notes": "State the fix version explicitly.",
        },
    ],
    "claims": CLAIMS,
}

VIDEO_PACKAGE = {
    "title": "Critical bypass explained",
    "scenes": [
        {
            "narration": "A critical vulnerability was disclosed this week.",
            "visual": "Title card with the product name",
            "on_screen_text": "CVSS 9.1",
            "duration_sec": 30,
            "source_chunks": ["c1"],
        },
        {
            "narration": "A fix is available now.",
            "visual": "Version callout",
            "on_screen_text": "Upgrade to 22.7R2.6",
            "duration_sec": 20,
            "source_chunks": ["c1"],
        },
    ],
    "storyboard_notes": "Two beats: severity, then action.",
    "subtitles_srt": "",
    "visual_recommendations": ["Avoid film stills - separately licensed."],
    "discussion_prompts": ["What is your patch window?"],
    "claims": CLAIMS,
}


# --- markdown / json --------------------------------------------------------


@pytest.mark.p0
def test_markdown_render_writes_a_readable_file(tmp_path):
    """TC-0701: file written, content matches the artefact JSON."""
    paths = dispatch.render(LINKEDIN, "markdown", str(tmp_path), basename="linkedin")
    md = next(p for p in paths if p.endswith(".md"))

    text = pathlib.Path(md).read_text(encoding="utf-8")
    assert "Upgrade to 22.7R2.6" in text
    assert "#CyberSecurity" in text
    # Citations are the point of the pipeline; they must survive to the file.
    assert "c1" in text


@pytest.mark.p1
def test_json_is_written_alongside_markdown(tmp_path):
    import json

    paths = dispatch.render(LINKEDIN, "markdown", str(tmp_path), basename="linkedin")
    js = next(p for p in paths if p.endswith(".json"))
    assert json.loads(pathlib.Path(js).read_text(encoding="utf-8")) == LINKEDIN


# --- docx / pdf -------------------------------------------------------------


@pytest.mark.p0
def test_docx_and_pdf_are_written_and_open(tmp_path):
    """TC-0701: the file opens and the content is there."""
    paths = dispatch.render(EXEC_SUMMARY, "docx", str(tmp_path), basename="summary")
    assert len(paths) == 2

    docx_path = next(p for p in paths if p.endswith(".docx"))
    pdf_path = next(p for p in paths if p.endswith(".pdf"))

    import docx

    doc = docx.Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Patch immediately." in text
    assert "Rated 9.1 on CVSS" in text

    # A PDF that opens starts with the magic number and has real length.
    with open(pdf_path, "rb") as fh:
        assert fh.read(5) == b"%PDF-"
    assert os.path.getsize(pdf_path) > 1000


# --- pptx -------------------------------------------------------------------


@pytest.mark.p1
def test_pptx_has_slides_and_populated_speaker_notes(tmp_path):
    """TC-0702: slides present AND speaker notes populated.

    A deck without notes is half an artefact.
    """
    from pptx import Presentation

    paths = dispatch.render(PRESENTATION, "pptx", str(tmp_path), basename="deck")
    prs = Presentation(paths[0])

    # Title slide + 2 content slides + sources slide.
    assert len(prs.slides) == 4

    notes = [
        s.notes_slide.notes_text_frame.text
        for s in prs.slides
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text
    ]
    assert any("Open by naming the product" in n for n in notes)
    assert any("State the fix version" in n for n in notes)


# --- video package ----------------------------------------------------------


@pytest.mark.p0
def test_no_generative_video_api_is_called(monkeypatch, tmp_path):
    """TC-0705, Invariant 10.

    Assert the export path makes no model call at all. Video is assembled from
    text the model already produced; rendering is TTS + ffmpeg.
    """
    from app.gateway import router

    calls = []

    async def spy(*a, **k):
        calls.append(a)
        return ""

    monkeypatch.setattr(router, "complete", spy)
    video.render_package(VIDEO_PACKAGE, str(tmp_path), "vid")
    assert calls == [], "the export path made a model call"


@pytest.mark.p1
def test_srt_cue_count_matches_scene_count(tmp_path):
    """TC-0703: cue count matches scenes; timings sum to the runtime."""
    srt = video.build_srt(VIDEO_PACKAGE["scenes"])

    cues = [b for b in srt.strip().split("\n\n") if b.strip()]
    assert len(cues) == len(VIDEO_PACKAGE["scenes"])

    # 30s + 20s = 50s total, so the last cue must end at 00:00:50,000.
    assert "00:00:50,000" in cues[-1]
    assert cues[0].startswith("1\n00:00:00,000 --> 00:00:30,000")


@pytest.mark.p1
def test_video_package_always_writes_package_and_subtitles(tmp_path):
    """The text package is the deliverable; media is a bonus on top."""
    paths = video.render_package(VIDEO_PACKAGE, str(tmp_path), "vid")

    assert any(p.endswith("vid.json") for p in paths)
    assert any(p.endswith("vid.srt") for p in paths)

    srt_path = next(p for p in paths if p.endswith(".srt"))
    srt = pathlib.Path(srt_path).read_text(encoding="utf-8")
    assert "-->" in srt


@pytest.mark.p1
def test_video_duration_matches_the_declared_scene_durations(tmp_path):
    """TC-0703/0704: the mp4 runs as long as duration_sec says.

    Regression: the first implementation used ffmpeg -shortest, so each scene
    lasted as long as its TTS narration instead of its declared duration. The
    SRT is timed from duration_sec, so the subtitles and the picture drifted
    apart - a 11s package rendered as a 6s video with subtitles running past
    the end.
    """
    import json
    import subprocess

    paths = video.render_package(VIDEO_PACKAGE, str(tmp_path), "vid")
    mp4 = [p for p in paths if p.endswith(".mp4")]
    if not mp4:
        pytest.skip("ffmpeg unavailable in this environment")

    probe = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "json", mp4[0]],
        capture_output=True,
        text=True,
        check=True,
    )
    actual = float(json.loads(probe.stdout)["format"]["duration"])
    declared = sum(s["duration_sec"] for s in VIDEO_PACKAGE["scenes"])

    assert abs(actual - declared) < 1.0, (
        f"mp4 runs {actual:.1f}s but the scenes declare {declared}s; "
        "subtitles will drift out of sync"
    )


@pytest.mark.p1
def test_title_cards_are_drawn_locally(tmp_path):
    """No stock imagery, no image generation - PIL draws them."""
    path = video.title_card(VIDEO_PACKAGE["scenes"][0], str(tmp_path / "card.png"))
    from PIL import Image

    with Image.open(path) as img:
        assert img.size == (video.WIDTH, video.HEIGHT)


# --- registry-driven --------------------------------------------------------


@pytest.mark.p0
def test_every_registered_renderer_is_implemented():
    """A registry entry naming an unimplemented renderer would fail at export.

    Checked here rather than at render time so the gap surfaces in CI, not
    mid-demo.
    """
    import inspect

    source = inspect.getsource(dispatch.render)
    for spec in registry.all_formats():
        assert f'"{spec.renderer}"' in source, f"{spec.id} names renderer {spec.renderer}"


@pytest.mark.p1
def test_unknown_renderer_raises_clearly(tmp_path):
    with pytest.raises(ValueError, match="Unknown renderer"):
        dispatch.render(LINKEDIN, "hologram", str(tmp_path))

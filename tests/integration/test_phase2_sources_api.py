"""Phase 2 done-criterion: POST /sources end to end against real services.

Requires the compose stack. Chunks must land in qdrant with a source_id payload.
"""

import io

import pytest
from fastapi.testclient import TestClient

from app.db.models import Source
from app.db.session import session_scope
from app.ingest import embed
from app.main import app

pytestmark = pytest.mark.integration

ADVISORY = (
    "Critical authentication bypass in NetGuard Connect Secure\n\n"
    "A vulnerability rated 9.1 out of 10 on the CVSS scale allows an "
    "unauthenticated attacker to bypass authentication entirely.\n\n"
    "The vendor has confirmed exploitation in the wild since 21 August 2026.\n\n"
    "A fix is available in version 22.7R2.6, released 28 August 2026.\n\n"
    "Approximately 14,000 internet-facing instances remain affected."
)


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def clean():
    """Remove test residue from both stores."""
    yield
    with session_scope() as s:
        for row in s.query(Source).all():
            if row.title.startswith(("Critical authentication", "probe", "Untitled")):
                embed.delete_source(row.id)
                s.delete(row)


def _pdf_bytes(text: str = ADVISORY) -> bytes:
    """A real single-page PDF, built in-process so no fixture file is needed."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox(fitz.Rect(40, 40, 550, 780), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.mark.p0
def test_text_pdf_ingests_and_embeds(client):
    """TC-0101: source_id returned, text extracted, chunks in qdrant."""
    r = client.post("/sources", files={"file": ("advisory.pdf", _pdf_bytes(), "application/pdf")})
    assert r.status_code == 201, r.text
    body = r.json()

    assert body["source_id"].startswith("s_")
    assert body["source_type"] == "pdf"
    assert body["chunk_count"] > 0
    assert body["reused"] is False

    # Chunks are retrievable and scoped by the source_id payload field.
    hits = embed.search([body["source_id"]], "what is the CVSS score", k=3)
    assert hits, "no chunks retrievable from qdrant"
    assert all(h["source_id"] == body["source_id"] for h in hits)


@pytest.mark.p0
def test_duplicate_upload_returns_existing_source(client):
    """TC-0109: second call returns the existing id; no re-extraction."""
    data = _pdf_bytes()
    first = client.post("/sources", files={"file": ("a.pdf", data, "application/pdf")})
    second = client.post("/sources", files={"file": ("a.pdf", data, "application/pdf")})

    assert first.status_code == 201
    assert second.status_code == 201
    assert second.json()["source_id"] == first.json()["source_id"]
    assert second.json()["reused"] is True


@pytest.mark.p0
def test_corrupt_pdf_fails_fast_with_a_clear_message(client):
    """TC-0110: fails in <3s with a message; no job, no tokens, no source row."""
    r = client.post(
        "/sources",
        files={"file": ("broken.pdf", b"%PDF-1.4 truncated garbage", "application/pdf")},
    )
    assert r.status_code == 400
    assert "PDF" in r.json()["detail"]

    with session_scope() as s:
        assert s.query(Source).filter(Source.title == "broken.pdf").count() == 0


@pytest.mark.p1
def test_unsupported_type_is_rejected(client):
    """TC-0113: rejected listing the accepted types."""
    r = client.post(
        "/sources", files={"file": ("x.exe", b"MZ\x90\x00", "application/octet-stream")}
    )
    assert r.status_code == 400
    assert ".pdf" in r.json()["detail"]


@pytest.mark.p0
def test_empty_source_is_rejected(client):
    """TC-0112: blank content -> no readable content."""
    r = client.post("/sources", data={"text": "   "})
    assert r.status_code == 400


@pytest.mark.p1
def test_free_form_text_source(client):
    """TC-0115: text with no file still creates a chunked source."""
    r = client.post("/sources", data={"text": ADVISORY})
    assert r.status_code == 201
    assert r.json()["chunk_count"] >= 1
    assert r.json()["source_type"] == "text"


@pytest.mark.p1
def test_docx_ingests(client):
    """TC-0103: headings preserved in the extracted text."""
    import docx as pydocx

    buf = io.BytesIO()
    doc = pydocx.Document()
    doc.add_heading("Critical authentication bypass", level=1)
    doc.add_paragraph(ADVISORY)
    doc.save(buf)

    r = client.post(
        "/sources",
        files={
            "file": (
                "a.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert r.status_code == 201, r.text
    assert r.json()["chunk_count"] >= 1


@pytest.mark.p1
def test_recent_sources_listing(client):
    """UC-12: the dashboard needs a recent-sources list."""
    client.post("/sources", data={"text": ADVISORY})
    r = client.get("/sources")
    assert r.status_code == 200
    assert isinstance(r.json(), list)


@pytest.mark.p0
def test_chunk_ids_survive_a_reload(client):
    """TC-0114: ids identical across reads; never renumbered."""
    from app.ingest.service import to_content_object

    r = client.post("/sources", files={"file": ("a.pdf", _pdf_bytes(), "application/pdf")})
    source_id = r.json()["source_id"]

    with session_scope() as s:
        row = s.get(Source, source_id)
        first = [c.id for c in to_content_object(row).chunks]
        second = [c.id for c in to_content_object(row).chunks]

    assert first == second
    assert first == [f"c{i}" for i in range(1, len(first) + 1)]

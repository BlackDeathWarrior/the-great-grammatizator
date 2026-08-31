"""docx, pdf and pptx rendering.

Generic over artefact shape for the same reason as markdown: the renderer is
named by the registry entry, so it must work for any format that names it.
"""

from __future__ import annotations

import logging

log = logging.getLogger(__name__)


def _title_of(artefact: dict) -> str:
    return artefact.get("title") or artefact.get("headline") or artefact.get("hook") or "Artefact"


def to_docx(artefact: dict, path: str) -> str:
    import docx

    doc = docx.Document()
    doc.add_heading(_title_of(artefact), level=0)

    for key, value in artefact.items():
        if key in ("claims", "title"):
            continue
        label = key.replace("_", " ").title()
        if isinstance(value, str) and value.strip():
            doc.add_heading(label, level=1)
            doc.add_paragraph(value)
        elif isinstance(value, list) and value:
            doc.add_heading(label, level=1)
            for item in value:
                if isinstance(item, str):
                    doc.add_paragraph(item, style="List Bullet")
                elif isinstance(item, dict):
                    for k, v in item.items():
                        doc.add_paragraph(f"{k.replace('_', ' ').title()}: {v}")

    claims = artefact.get("claims") or []
    if claims:
        doc.add_heading("Sources", level=1)
        for claim in claims:
            doc.add_paragraph(
                f"“{claim.get('text', '')}” — {claim.get('chunk_id', '')}",
                style="List Bullet",
            )

    doc.save(path)
    return path


def to_pdf(artefact: dict, path: str) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(_escape(_title_of(artefact)), styles["Title"]), Spacer(1, 12)]

    for key, value in artefact.items():
        if key in ("claims", "title"):
            continue
        label = key.replace("_", " ").title()
        if isinstance(value, str) and value.strip():
            story += [
                Paragraph(_escape(label), styles["Heading2"]),
                Paragraph(_escape(value), styles["BodyText"]),
                Spacer(1, 8),
            ]
        elif isinstance(value, list) and value:
            story.append(Paragraph(_escape(label), styles["Heading2"]))
            items = []
            for item in value:
                text = (
                    item
                    if isinstance(item, str)
                    else "; ".join(f"{k}: {v}" for k, v in item.items())
                )
                items.append(ListItem(Paragraph(_escape(str(text)), styles["BodyText"])))
            story += [ListFlowable(items, bulletType="bullet"), Spacer(1, 8)]

    claims = artefact.get("claims") or []
    if claims:
        story.append(Paragraph("Sources", styles["Heading2"]))
        story.append(
            ListFlowable(
                [
                    ListItem(
                        Paragraph(
                            _escape(f"“{c.get('text', '')}” - {c.get('chunk_id', '')}"),
                            styles["BodyText"],
                        )
                    )
                    for c in claims
                ],
                bulletType="bullet",
            )
        )

    SimpleDocTemplate(path, pagesize=A4).build(story)
    return path


def to_pptx(artefact: dict, path: str) -> str:
    """Slides plus SPEAKER NOTES (TC-0702).

    Notes are the half a deck is useless without, so they are populated from
    the schema's speaker_notes rather than left empty.
    """
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    opening = prs.slides.add_slide(title_layout)
    opening.shapes.title.text = _title_of(artefact)
    if len(opening.placeholders) > 1:
        opening.placeholders[1].text = artefact.get("bottom_line", "")

    bullet_layout = prs.slide_layouts[1]
    for slide_spec in artefact.get("slides") or []:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_spec.get("heading", "")

        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide_spec.get("bullets") or []):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
            para.font.size = Pt(18)

        notes = slide_spec.get("speaker_notes", "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    claims = artefact.get("claims") or []
    if claims:
        sources = prs.slides.add_slide(bullet_layout)
        sources.shapes.title.text = "Sources"
        frame = sources.placeholders[1].text_frame
        frame.clear()
        for i, claim in enumerate(claims):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = f"“{claim.get('text', '')}” - {claim.get('chunk_id', '')}"
            para.font.size = Pt(14)

    prs.save(path)
    return path


def _escape(text: str) -> str:
    """reportlab's Paragraph parses a mini-HTML dialect."""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

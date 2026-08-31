"""DOCX extraction via python-docx. Headings preserved in the text (TC-0103)."""

from __future__ import annotations

import io

from app.ingest.errors import CorruptSource


def extract(data: bytes) -> tuple[str, str]:
    """Return (text, title)."""
    import docx as pydocx

    try:
        doc = pydocx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise CorruptSource(f"Could not open DOCX: {exc}") from exc

    parts: list[str] = []
    title = ""
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            # Keep headings visible as structure so chunking splits on them and
            # the generator can see document shape.
            parts.append(f"\n## {text}\n")
            if not title and style in ("heading 1", "title"):
                title = text
        else:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip(), title
